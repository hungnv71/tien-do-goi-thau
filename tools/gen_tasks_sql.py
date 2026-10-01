"""Sinh supabase/05_import_nhiem_vu_tuan40.sql từ BC_TienDo_PhongQLHT_Tuan40_2026.xlsx (sheet ChiTiet_ToanPhong).

Chạy: python tools/gen_tasks_sql.py
- Cập nhật họ tên thật + email cho cán bộ; thêm cán bộ mới; thêm tài khoản Trưởng phòng thainn2 (chưa có PIN).
- Nhập 76 nhiệm vụ: % (0..1) -> 0..100; hạn giao = hạn hiện hành = hạn trong file; hạn KHÓA.
- Không suy diễn: ngày hoàn thành thực tế để trống (file không có).
- ID ổn định theo tài khoản + tên + hạn => chạy lại không nhân đôi.
"""
import datetime as dt
import hashlib
import pathlib
import re
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
SRC = ROOT / "BC_TienDo_PhongQLHT_Tuan40_2026.xlsx"
OUT = ROOT / "supabase" / "05_import_nhiem_vu_tuan40.sql"
WEEK, WEEK_DAY = "2026-W40", "2026-09-29"

# tài khoản -> (họ tên, email, vị trí). Email theo mẫu tài khoản@viettel.com.vn; để trống nếu tài khoản không theo mẫu.
STAFF = {
    "thainn2": ("Thái (Trưởng phòng)", "thainn2@viettel.com.vn", 0),
    "hungnv71": ("Nguyễn Việt Hùng", "hungnv71@viettel.com.vn", 1),
    "hungnt16": ("Nguyễn Thanh Hùng", "hungnt16@viettel.com.vn", 2),
    "dangnt2": ("Nguyễn Thị Dàng", "dangnt2@viettel.com.vn", 3),
    "tungtt17": ("Trần Thanh Tùng", "tungtt17@viettel.com.vn", 4),
    "tucna": ("Nguyễn Anh Túc", "tucna@viettel.com.vn", 5),
    "bichngoc": ("Nguyễn Thị Bích Ngọc", None, 6),
    "doanhtv4": ("Trân Văn Doanh", "doanhtv4@viettel.com.vn", 7),   # giữ đúng chính tả trong file nguồn
    "vietpt098": ("Phạm Trọng Viết", "vietpt098@viettel.com.vn", 8),
    "hoanglehai2605": ("Hoàng Lê Hải", None, 9),
    "tuantn1": ("Tạ Ngọc Tuấn", "tuantn1@viettel.com.vn", 10),
}
STATUS = {"hoàn thành": "done", "đang thực hiện": "in_progress", "chưa bắt đầu": "not_started",
          "chậm tiến độ": "in_progress", "chờ ý kiến / phối hợp": "waiting"}


def q(v):
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return repr(round(v, 2)) if isinstance(v, float) else str(v)
    if isinstance(v, (dt.date, dt.datetime)):
        return "'" + v.strftime("%Y-%m-%d") + "'"
    return "'" + str(v).replace("'", "''") + "'"


def clean(s, keep_lines=True):
    if s is None:
        return None
    s = str(s).replace("\r", "")
    s = "\n".join(re.sub(r"[ \t]+", " ", x).strip() for x in s.split("\n")) if keep_lines else re.sub(r"\s+", " ", s)
    s = s.strip()
    return s or None


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["ChiTiet_ToanPhong"]
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0] is not None and r[4]]
    # Ai đã gửi báo cáo tuần: theo sheet DASHBOARD (cột "File nguồn"; "Chưa gửi BC tuần" = chưa gửi)
    name2acct = {clean(r[1], False): (clean(r[2]) or "").lower() for r in rows}
    sub = set()
    for d in wb["DASHBOARD"].iter_rows(min_row=31, values_only=True):
        nm, src = clean(d[1], False) if len(d) > 1 else None, " ".join(str(x) for x in d[2:] if x is not None)
        if nm in name2acct and "chưa gửi" not in src.lower():
            sub.add(name2acct[nm])

    L = ["-- =====================================================================",
         "-- Migration 05: NHẬP NHIỆM VỤ PHÒNG tuần 40/2026 từ BC_TienDo_PhongQLHT_Tuan40_2026.xlsx",
         f"-- Sinh tự động {dt.date.today():%d/%m/%Y} bởi tools/gen_tasks_sql.py — {len(rows)} nhiệm vụ. Chạy lại an toàn.",
         "-- =====================================================================", ""]
    # cán bộ
    for acct, (name, email, pos) in STAFF.items():
        role = "manager" if acct == "thainn2" else "staff"
        L.append(f"insert into staff (id, full_name, account, email, role, is_manager, position) "
                 f"select {q('st_' + acct)}, {q(name)}, {q(acct)}, {q(email)}, {q(role)}, {q(role == 'manager')}, {pos} "
                 f"where not exists (select 1 from staff where lower(account) = {q(acct)});")
        L.append(f"update staff set full_name = {q(name)}, email = coalesce(email, {q(email)}), position = {pos} where lower(account) = {q(acct)};")
    L.append("update staff set role = 'admin' where lower(account) = 'hungnv71';")
    L.append("-- Tài khoản 'Trưởng phòng' chung được thay bằng tài khoản thật thainn2 (giữ lại, chỉ ngừng đăng nhập)")
    L.append("update staff set active = false where id = 'st_truongphong' and not exists (select 1 from packages where staff_id = 'st_truongphong') and not exists (select 1 from contracts where staff_id = 'st_truongphong');")
    L.append("")

    vals = []
    for i, r in enumerate(rows, 1):
        acct = (clean(r[2]) or "").lower()
        title = clean(r[4])
        typ = clean(r[3], False) or "Nhiệm vụ được giao"
        src_doc = None
        if typ.lower().startswith("phiếu giao nhiệm vụ"):
            src_doc, typ = typ, "Phiếu giao nhiệm vụ quý"
        due = r[6] if isinstance(r[6], (dt.date, dt.datetime)) else None
        giao = r[5] if isinstance(r[5], (dt.date, dt.datetime)) else None
        pct = r[8] if isinstance(r[8], (int, float)) else None
        pct = round(pct * 100, 1) if pct is not None and pct <= 1.0001 else pct
        st = STATUS.get((clean(r[9], False) or "").lower(), "in_progress")
        wk_res = clean(r[12])
        if wk_res and "chưa gửi" in wk_res.lower():
            wk_res = None
        tid = "nv_" + hashlib.md5(f"{acct}|{title}|{due}".encode("utf-8")).hexdigest()[:12]
        pkg = "'pk_demo2'" if title and "3184 trạm 5G" in title else "null"
        note = clean(r[11]) if False else None
        vals.append("(" + ",".join([
            q(tid), q(f"NV26-{i:04d}"), q(title), q(typ), q(src_doc), pkg, q(acct), q(giao), q(due), q(due),
            q(clean(r[7])), q(pct), q(st), q(clean(r[11])), q(wk_res), q(clean(r[13])), q(clean(r[14])), q(clean(r[15])),
            q(typ == "Nhiệm vụ thường xuyên"), q(WEEK if acct in sub else None), q(WEEK_DAY if acct in sub else None),
            q("Bảng ChiTiet_ToanPhong tuần 40/2026" + (" · " + clean(r[17], False) if r[17] else "")),
        ]) + ")")
    L.append("with v(id,code,title,typ,src,pkg,acct,giao,due,odue,outp,pct,st,res,wk,nxt,diff,prop,rec,lw,lwd,source) as (values")
    L.append(",\n".join(vals))
    L.append(""") insert into tasks (id, code, title, task_type, source_doc, package_id, staff_id, assigned_date, due_date, original_due, due_locked,
  output, percent, status, result_total, week_result, next_plan, difficulty, proposal, recurring, last_report_week, last_report_at, source)
select v.id, v.code, v.title, v.typ, v.src, v.pkg, (select s.id from staff s where lower(s.account) = v.acct limit 1),
  v.giao::date, v.due::date, v.odue::date, true, v.outp, v.pct::numeric, v.st, v.res, v.wk, v.nxt, v.diff, v.prop, v.rec::boolean,
  v.lw, v.lwd::timestamptz, v.source
from v on conflict (id) do nothing;
""")
    for acct in sorted(sub):
        n = sum(1 for r in rows if (clean(r[2]) or "").lower() == acct)
        L.append(f"insert into weekly_reports (id, staff_id, week, submitted_at, task_count, note) "
                 f"select s.id || '|{WEEK}', s.id, '{WEEK}', '{WEEK_DAY}', {n}, 'Nhập từ file tổng hợp tuần 40' from staff s where lower(s.account) = {q(acct)} "
                 f"on conflict (staff_id, week) do nothing;")
    L.append("")
    L.append("select (select count(*) from tasks) nhiem_vu, (select count(*) from tasks where staff_id is null) thieu_cb, (select count(*) from weekly_reports) bc_tuan;")
    OUT.write_text("\n".join(L), encoding="utf-8")
    print(f"OK: {len(rows)} nhiệm vụ, {len(sub)} cán bộ đã báo cáo tuần -> {OUT}")


if __name__ == "__main__":
    main()
